import os
import re

import pdfplumber
from docx import Document

import config
from src.ocr import run_ocr_on_image, run_ocr_on_pdf


def check_extension(filename: str) -> bool:
    ext = os.path.splitext(filename)[1].lower()
    return ext in config.ALLOWED_EXTENSIONS


def check_file_size(file_path: str) -> tuple[bool, str]:
    size_bytes = os.path.getsize(file_path)
    if size_bytes == 0:
        return False, "The uploaded file is empty."
    max_bytes = config.MAX_FILE_SIZE_MB * 1024 * 1024
    if size_bytes > max_bytes:
        return False, f"File exceeds the {config.MAX_FILE_SIZE_MB}MB limit."
    return True, ""


def check_pdf_page_count(file_path: str) -> tuple[bool, str]:
    try:
        with pdfplumber.open(file_path) as pdf:
            if len(pdf.pages) > config.MAX_PDF_PAGES:
                return False, f"PDF exceeds the {config.MAX_PDF_PAGES}-page limit."
        return True, ""
    except Exception:
        return False, "Could not open PDF - file may be corrupted."


def validate_upload(file_path: str, original_filename: str) -> tuple[bool, str]:
    if not check_extension(original_filename):
        return False, f"Unsupported file type. Allowed: {', '.join(sorted(config.ALLOWED_EXTENSIONS))}"
    size_ok, size_msg = check_file_size(file_path)
    if not size_ok:
        return False, size_msg

    ext = os.path.splitext(original_filename)[1].lower()
    if ext == '.pdf':
        page_ok, page_msg = check_pdf_page_count(file_path)
        if not page_ok:
            return False, page_msg
    return True, "File is valid."


def extract_pdf_text(file_path: str) -> str:
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            if i >= config.MAX_PDF_PAGES:
                break
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_docx_text(file_path: str) -> str:
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def is_text_sufficient(text: str, min_chars: int = 20) -> bool:
    return len(text.strip()) >= min_chars


def load_document(file_path: str, original_filename: str) -> str:
    is_valid, message = validate_upload(file_path, original_filename)
    if not is_valid:
        raise ValueError(message)

    ext = os.path.splitext(original_filename)[1].lower()

    if ext == '.pdf':
        text = extract_pdf_text(file_path)
        if not is_text_sufficient(text):
            text = run_ocr_on_pdf(file_path, max_pages=config.MAX_PDF_PAGES)
    elif ext == '.docx':
        text = extract_docx_text(file_path)
    elif ext in ('.jpg', '.jpeg', '.png'):
        text = run_ocr_on_image(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not is_text_sufficient(text):
        raise ValueError("Could not extract readable text from this document.")

    text = clean_extracted_text(text)

    return text


def clean_extracted_text(text: str) -> str:
    text = text.replace('\x00', '')
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()
