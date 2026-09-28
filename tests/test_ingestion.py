import os

import pytest
from PIL import Image
from reportlab.pdfgen import canvas

import config
from src import document_loader as dl
from conftest import SAMPLES


def make_pdf(path, pages, text="Hello world this is page text"):
    c = canvas.Canvas(str(path))
    for i in range(pages):
        c.drawString(72, 700, f"{text} {i}")
        c.showPage()
    c.save()
    return str(path)


def test_rejects_unsupported_type(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("hello")
    with pytest.raises(ValueError, match="Unsupported"):
        dl.load_document(str(f), "notes.txt")


def test_rejects_empty_file(tmp_path):
    f = tmp_path / "empty.pdf"
    f.write_bytes(b"")
    with pytest.raises(ValueError, match="empty"):
        dl.load_document(str(f), "empty.pdf")


def test_rejects_corrupted_pdf(tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"this is not a pdf at all")
    with pytest.raises(ValueError, match="corrupted"):
        dl.load_document(str(f), "bad.pdf")


def test_rejects_oversized_file(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MAX_FILE_SIZE_MB", 0)
    f = tmp_path / "big.pdf"
    f.write_bytes(b"%PDF-1.4 x")
    with pytest.raises(ValueError, match="limit"):
        dl.load_document(str(f), "big.pdf")


def test_rejects_too_many_pages(tmp_path):
    path = make_pdf(tmp_path / "long.pdf", config.MAX_PDF_PAGES + 1)
    with pytest.raises(ValueError, match="page limit"):
        dl.load_document(path, "long.pdf")


def test_text_pdf_extracts_text():
    text = dl.load_document(os.path.join(SAMPLES, "sample_invoice.pdf"), "sample_invoice.pdf")
    assert "INV-2026-0042" in text


def test_docx_extracts_text():
    text = dl.load_document(os.path.join(SAMPLES, "sample_contract.docx"), "sample_contract.docx")
    assert "Delta Software" in text


def test_docx_extracts_table_text(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Invoice details")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Widget"
    table.cell(1, 1).text = "42.00"
    path = tmp_path / "table.docx"
    doc.save(str(path))
    text = dl.load_document(str(path), "table.docx")
    assert "Widget" in text and "42.00" in text


def test_image_uses_ocr():
    text = dl.load_document(os.path.join(SAMPLES, "sample_scan.png"), "sample_scan.png")
    assert "Green Market" in text


def test_scanned_pdf_falls_back_to_ocr(tmp_path):
    out = tmp_path / "scan.pdf"
    Image.open(os.path.join(SAMPLES, "sample_scan.png")).convert("RGB").save(out, "PDF", resolution=150)
    text = dl.load_document(str(out), "scan.pdf")
    assert "Green Market" in text


def test_blank_image_is_not_treated_as_a_document(tmp_path):
    blank = tmp_path / "blank.png"
    Image.new("RGB", (300, 300), "white").save(blank)
    with pytest.raises(ValueError, match="readable text"):
        dl.load_document(str(blank), "blank.png")
